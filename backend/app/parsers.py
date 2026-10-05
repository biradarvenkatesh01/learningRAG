from pathlib import Path

from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def parse_pdf(path: Path) -> list[dict]:
    reader = PdfReader(path)
    sections = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            sections.append({"text": text, "source": f"page {i + 1}"})
    return sections


def parse_docx(path: Path) -> list[dict]:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    # Word tables aren't paragraphs, so grab them separately
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts)
    return [{"text": text, "source": "document"}] if text.strip() else []


def parse_pptx(path: Path) -> list[dict]:
    prs = Presentation(path)
    sections = []
    for i, slide in enumerate(prs.slides):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
        text = "\n".join(parts)
        if text.strip():
            sections.append({"text": text, "source": f"slide {i + 1}"})
    return sections


def parse_txt(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [{"text": text, "source": "document"}] if text.strip() else []


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".pptx": parse_pptx,
    ".txt": parse_txt,
    ".md": parse_txt,
}


def parse_document(path: str | Path) -> list[dict]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in PARSERS:
        raise ValueError(f"Unsupported file type: {ext}. Supported: {list(PARSERS)}")
    return PARSERS[ext](path)